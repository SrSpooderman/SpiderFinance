"""Lectura acotada de CSV/XLSX y normalización de filas financieras."""

import csv
import io
import json
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 50
PORTABLE_FIELDS = [
    "date", "type", "concept", "amount", "source_account", "destination_account",
    "category", "status", "notes", "payment_method", "is_fixed", "is_necessary",
]


def string_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def unique_headers(values: list[str]) -> list[str]:
    result = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, 1):
        base = value.strip() or f"Columna {index}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
    return result


def sheet_rows(rows, name: str) -> tuple[list[str], list[tuple[int, dict]]]:
    iterator = iter(rows)
    try:
        headers = unique_headers([string_value(value) for value in next(iterator)])
    except StopIteration:
        raise ValueError(f"La hoja {name} está vacía") from None
    if not headers or len(headers) > MAX_COLUMNS:
        raise ValueError("Demasiadas columnas")
    data = []
    for row_number, values in enumerate(iterator, 2):
        if len(data) >= MAX_ROWS:
            raise ValueError(f"El archivo supera {MAX_ROWS} filas")
        cells = [string_value(value) for value in values]
        if any(cells):
            data.append((row_number, dict(zip(headers, cells, strict=False))))
    return headers, data


def parse_file(contents: bytes, filename: str) -> dict[str, tuple[list[str], list[tuple[int, dict]]]]:
    if len(contents) > MAX_FILE_BYTES:
        raise ValueError("El archivo supera 5 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        try:
            records = json.loads(contents.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise ValueError("El JSON no es válido o no está codificado en UTF-8") from None
        if not isinstance(records, list) or any(not isinstance(record, dict) for record in records):
            raise ValueError("El JSON debe ser una lista de movimientos")
        if len(records) > MAX_ROWS:
            raise ValueError(f"El archivo supera {MAX_ROWS} filas")
        headers = list(dict.fromkeys(key for record in records for key in record)) if records else PORTABLE_FIELDS
        if not headers or len(headers) > MAX_COLUMNS or any(not isinstance(key, str) for key in headers):
            raise ValueError("El JSON no contiene columnas válidas")
        return {"JSON": sheet_rows([headers, *([record.get(key) for key in headers] for record in records)], "JSON")}
    if suffix == ".csv":
        try:
            content = contents.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("El CSV debe estar codificado en UTF-8") from None
        sample = content[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return {"CSV": sheet_rows(csv.reader(io.StringIO(content), dialect), "CSV")}
    if suffix != ".xlsx":
        raise ValueError("Solo se admiten CSV, XLSX y JSON")
    try:
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("El XLSX descomprimido supera 20 MB")
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True, keep_links=False)
    except (zipfile.BadZipFile, InvalidFileException, OSError, KeyError, TypeError) as exc:
        raise ValueError("El XLSX no es válido") from exc
    try:
        if not workbook.sheetnames:
            raise ValueError("El XLSX no contiene hojas")
        result = {}
        total = 0
        for sheet in workbook:
            headers, rows = sheet_rows(sheet.values, sheet.title)
            total += len(rows)
            if total > MAX_ROWS:
                raise ValueError(f"El archivo supera {MAX_ROWS} filas")
            result[sheet.title] = (headers, rows)
        return result
    finally:
        workbook.close()


def parse_date(value: str) -> date:
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha inválida: {value}")


def parse_amount(value: str) -> Decimal:
    raw = value.strip().replace(" ", "").replace("€", "").replace("$", "")
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
    elif "," in raw:
        raw = raw.replace(",", ".") if len(raw.rsplit(",", 1)[1]) <= 2 else raw.replace(",", "")
    elif "." in raw and len(raw.rsplit(".", 1)[1]) == 3:
        raw = raw.replace(".", "")
    try:
        amount = Decimal(raw)
    except InvalidOperation:
        raise ValueError(f"Importe inválido: {value}") from None
    if not amount.is_finite() or amount == 0 or amount.as_tuple().exponent < -2:
        raise ValueError(f"Importe inválido: {value}")
    return amount


def parse_bool(value: str) -> bool:
    normalized = value.strip().casefold()
    if normalized in ("", "false", "0", "no"):
        return False
    if normalized in ("true", "1", "yes", "si", "sí"):
        return True
    raise ValueError(f"Valor booleano inválido: {value}")


def unescape_spreadsheet(value: str) -> str:
    return value[1:] if value.startswith(("'=", "'+", "'-", "'@")) else value


def parse_row(raw: dict, mapping: dict[str, str], default_account_id: int | None,
              positive_is_income: bool, spreadsheet_safe: bool = False) -> dict:
    def field(name: str) -> str:
        value = string_value(raw.get(mapping.get(name, ""), ""))
        return unescape_spreadsheet(value) if spreadsheet_safe else value

    day = parse_date(field("date"))
    amount = parse_amount(field("amount"))
    kind = field("type").upper()
    aliases = {"INGRESO": "INCOME", "GASTO": "EXPENSE", "TRANSFERENCIA": "TRANSFER",
               "INCOME": "INCOME", "EXPENSE": "EXPENSE", "TRANSFER": "TRANSFER"}
    if kind:
        kind = aliases.get(kind, kind)
    else:
        kind = "INCOME" if (amount > 0) == positive_is_income else "EXPENSE"
    concept = field("concept")
    if not concept:
        raise ValueError("Falta el concepto")
    source = field("source_account_id")
    destination = field("destination_account_id")
    account = field("account_id")
    selected = int(account) if account else default_account_id
    source_id = int(source) if source else (selected if kind == "EXPENSE" else None)
    destination_id = int(destination) if destination else (selected if kind == "INCOME" else None)
    category = field("category_id")
    status = field("status").upper() or "CLEARED"
    notes = field("notes")
    payment_method = field("payment_method")
    return {
        "date": day.isoformat(), "type": kind, "source_account_id": source_id,
        "destination_account_id": destination_id, "category_id": int(category) if category else None,
        "concept": concept, "amount": str(abs(amount)), "status": status,
        "notes": notes or None, "payment_method": payment_method or None,
        "is_fixed": parse_bool(field("is_fixed")), "is_necessary": parse_bool(field("is_necessary")),
    }
