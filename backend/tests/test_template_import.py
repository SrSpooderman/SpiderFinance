from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.auth import DEFAULT_CATEGORIES
from app.application.template_import import import_template, read_template
from app.core.db import Base
from app.infrastructure.models import (
    Account, Budget, Category, GoalContribution, IncomeReceipt, IncomeSource,
    RecurringExpense, RecurringPayment, Reservation, SavingsGoal, SavingsRule,
    ScheduledExpense, Transaction, User,
)


def workbook_file(tmp_path, expected_personal=Decimal("84.00")):
    workbook = Workbook()
    accounts = workbook.active
    accounts.title = "Cuentas"
    accounts.append(["Cuenta", "Tipo", "Banco/Entidad", "Saldo inicial", "Saldo actual", "Activa", "Notas"])
    accounts.append(["Personal", "Corriente", "BBVA", 0, float(expected_personal), "Sí", "Cuenta real"])
    accounts.append(["Ahorros", "Ahorro", "Virtual", 0, 10, "Sí", "Separación virtual"])
    config = workbook.create_sheet("Configuración")
    config.append(["Categorías"])
    for name in DEFAULT_CATEGORIES:
        config.append([name])
    movements = workbook.create_sheet("Movimientos")
    movements.append(["Fecha", "Cuenta origen", "Cuenta destino", "Tipo", "Categoría", "Subcategoría", "Concepto", "Importe", "Fijo", "Necesario", "Método de pago", "Notas"])
    movements.append([datetime(2026, 9, 30), None, "Personal", "Ingreso", "Otros", None, "Nómina septiembre", 100, "Sí", "Sí", "Transferencia", "Pagado"])
    movements.append([datetime(2026, 9, 30), "Personal", None, "Gasto", "Suscripciones", None, "Spotify", 5, "Sí", "No", "Tarjeta", "Pagado"])
    movements.append([datetime(2026, 9, 30), "Personal", "Ahorros", "Transferencia", "Otros", "Ahorro mensual", "Ahorro 10% septiembre", 10, "Sí", "Sí", "Transferencia", "Virtual"])
    movements.append([datetime(2026, 10, 1), "Personal", None, "Gasto", "Alimentación", "Merienda", "Merienda", 1, "No", "Sí", "Tarjeta", "Pagado"])
    budgets = workbook.create_sheet("Presupuesto")
    budgets.append(["Categoría", "Presupuesto mensual"])
    budgets.append(["Alimentación", 50])
    fixed = workbook.create_sheet("Gastos fijos")
    fixed.append(["Concepto", "Categoría", "Importe", "Periodicidad", "Día aproximado", "Cuenta", "Activo", "Inicio", "Fin"])
    fixed.append(["Spotify", "Suscripciones", 5, "Mensual", 1, "Personal", "Sí", None, None])
    fixed.append(["Compra", "Tecnología", 25, "Único", 15, "Personal", "Sí", datetime(2026, 10, 15), datetime(2026, 10, 15)])
    income = workbook.create_sheet("Ingresos fijos")
    income.append(["Concepto", "Importe", "Periodicidad", "Cuenta destino", "Activo", "Notas"])
    income.append(["Nómina", 100, "Mensual", "Personal", "Sí", "Cobro el último día del mes"])
    goals = workbook.create_sheet("Objetivos")
    goals.append(["Objetivo", "Objetivo €", "Ahorrado", "Fecha objetivo"])
    goals.append(["Viaje", 100, 10, datetime(2026, 12, 31)])
    path = tmp_path / "plantilla.xlsx"
    workbook.save(path)
    return path


def test_template_imports_entities_preserves_balances_and_refuses_duplicate(tmp_path):
    data = read_template(workbook_file(tmp_path))
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    try:
        with Session(engine) as db:
            user = User(email="import@example.com", password_hash="unused")
            db.add(user)
            db.flush()
            db.add_all(Category(user_id=user.id, name=name) for name in DEFAULT_CATEGORIES)
            db.commit()
            counts = import_template(db, user.id, data)
            db.commit()
            assert counts["transactions"] == 4
            assert counts["categories"] == len(DEFAULT_CATEGORIES) + 2
            accounts = {item.name: item for item in db.scalars(select(Account).where(Account.user_id == user.id))}
            assert accounts["Personal"].initial_balance == 0
            assert db.scalar(select(func.count()).select_from(Transaction)) == 4
            assert db.scalar(select(func.count()).select_from(Budget)) == 1
            assert db.scalar(select(func.count()).select_from(RecurringExpense)) == 1
            assert db.scalar(select(func.count()).select_from(ScheduledExpense)) == 1
            assert db.scalar(select(func.count()).select_from(IncomeSource)) == 1
            assert db.scalar(select(func.count()).select_from(IncomeReceipt)) == 1
            assert db.scalar(select(func.count()).select_from(RecurringPayment)) == 1
            assert db.scalar(select(func.count()).select_from(SavingsRule)) == 1
            assert db.scalar(select(SavingsRule.value)) == Decimal("10")
            assert db.scalar(select(Reservation.amount)) == Decimal("10")
            assert db.scalar(select(GoalContribution.amount)) == Decimal("10")
            assert db.scalar(select(SavingsGoal.target_amount)) == Decimal("100")
            assert db.scalar(select(Account.notes).where(Account.name == "Ahorros")) == "Separación virtual"
            with pytest.raises(ValueError, match="cuenta vacía"):
                import_template(db, user.id, data)
            db.rollback()
    finally:
        engine.dispose()


def test_template_rejects_mismatched_account_balance(tmp_path):
    with pytest.raises(ValueError, match="saldo calculado"):
        read_template(workbook_file(tmp_path, expected_personal=Decimal("99.00")))
