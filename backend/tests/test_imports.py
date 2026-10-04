from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest


def test_csv_preview_confirm_dedup_and_exports(client, auth):
    account = client.post("/api/v1/accounts", headers=auth, json={
        "name": "Banco", "type": "CHECKING", "currency": "EUR",
    }).json()["id"]
    today = date.today().isoformat()
    content = (
        f"Fecha;Concepto;Importe\n{today};Nómina;100,00\n"
        f"{today};Alquiler;-25,00\n{today};Fila inválida;abc\n"
    ).encode()
    uploaded = client.post("/api/v1/imports", headers=auth, files={
        "file": ("extracto.csv", content, "text/csv"),
    })
    assert uploaded.status_code == 201, uploaded.text
    job_id = uploaded.json()["id"]
    payload = {
        "sheet_name": "CSV",
        "mapping": {"date": "Fecha", "concept": "Concepto", "amount": "Importe"},
        "default_account_id": account,
        "positive_is_income": True,
    }
    preview = client.post(f"/api/v1/imports/{job_id}/preview", headers=auth, json=payload)
    assert preview.status_code == 200, preview.text
    assert [row["status"] for row in preview.json()["rows"]] == ["READY", "READY", "ERROR"]
    confirmed = client.post(f"/api/v1/imports/{job_id}/confirm", headers=auth)
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["imported_rows"] == 2
    assert confirmed.json()["skipped_rows"] == 1
    assert client.get(f"/api/v1/accounts/{account}/balance", headers=auth).json()["balance"] == "75.00"
    movement_id = client.get("/api/v1/transactions", headers=auth).json()["items"][0]["id"]
    assert client.delete(f"/api/v1/transactions/{movement_id}", headers=auth).status_code == 409
    duplicate = client.post("/api/v1/imports", headers=auth, files={
        "file": ("extracto.csv", content, "text/csv"),
    }).json()["id"]
    second = client.post(f"/api/v1/imports/{duplicate}/preview", headers=auth, json=payload)
    assert [row["status"] for row in second.json()["rows"]] == ["DUPLICATE", "DUPLICATE", "ERROR"]
    assert client.delete(f"/api/v1/imports/{job_id}", headers=auth).status_code == 204
    assert client.get("/api/v1/transactions", headers=auth).json()["total"] == 2

    exported_json = client.get("/api/v1/exports/transactions?format=json", headers=auth)
    assert exported_json.status_code == 200
    assert len(exported_json.json()) == 2
    exported_csv = client.get("/api/v1/exports/transactions?format=csv", headers=auth)
    assert exported_csv.status_code == 200
    assert "Nómina" in exported_csv.text
    exported_xlsx = client.get("/api/v1/exports/transactions?format=xlsx", headers=auth)
    assert exported_xlsx.status_code == 200
    workbook = load_workbook(BytesIO(exported_xlsx.content), read_only=True)
    assert len(list(workbook.active.values)) == 3
    workbook.close()


def test_xlsx_sheet_selection_and_isolation(client, auth):
    account = client.post("/api/v1/accounts", headers=auth, json={
        "name": "Banco", "type": "CHECKING", "currency": "EUR",
    }).json()["id"]
    workbook = Workbook()
    first = workbook.active
    first.title = "Resumen"
    first.append(["Dato"])
    first.append(["Ignorar"])
    second = workbook.create_sheet("Movimientos")
    second.append(["Fecha", "Concepto", "Importe"])
    second.append([date.today(), "Compra", -12.5])
    buffer = BytesIO()
    workbook.save(buffer)
    upload = client.post("/api/v1/imports", headers=auth, files={
        "file": ("banco.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    })
    assert upload.status_code == 201, upload.text
    job_id = upload.json()["id"]
    assert upload.json()["sheet_names"] == ["Resumen", "Movimientos"]
    preview = client.post(f"/api/v1/imports/{job_id}/preview", headers=auth, json={
        "sheet_name": "Movimientos",
        "mapping": {"date": "Fecha", "concept": "Concepto", "amount": "Importe"},
        "default_account_id": account,
    })
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows"][0]["status"] == "READY"
    other = client.post("/api/v1/auth/register", json={
        "email": "imports-other@example.com", "password": "another-strong-password",
    })
    other_auth = {"Authorization": f"Bearer {other.json()['access_token']}"}
    assert client.get(f"/api/v1/imports/{job_id}", headers=other_auth).status_code == 404


@pytest.mark.parametrize("format", ["csv", "xlsx", "json"])
def test_export_import_round_trip_across_instances(client, auth, format):
    def account(headers, name):
        response = client.post("/api/v1/accounts", headers=headers, json={
            "name": name, "type": "CHECKING", "currency": "EUR",
        })
        assert response.status_code == 201, response.text
        return response.json()["id"]

    def category(headers, name, parent_id=None):
        response = client.post("/api/v1/categories", headers=headers, json={
            "name": name, "parent_id": parent_id,
        })
        assert response.status_code == 201, response.text
        return response.json()["id"]

    source_bank = account(auth, "Banco")
    source_cash = account(auth, "Efectivo")
    source_home = category(auth, "Casa personal")
    source_rent = category(auth, "Alquiler", source_home)
    today = date.today().isoformat()
    movements = [
        {"type": "INCOME", "destination_account_id": source_bank, "concept": "Nómina", "amount": "1000.00"},
        {"type": "EXPENSE", "source_account_id": source_bank, "category_id": source_rent,
         "concept": "=Alquiler", "amount": "300.00", "status": "PENDING",
         "is_fixed": True, "is_necessary": True, "notes": "+nota"},
        {"type": "TRANSFER", "source_account_id": source_bank, "destination_account_id": source_cash,
         "concept": "Traspaso", "amount": "100.00"},
        {"type": "ADJUSTMENT", "source_account_id": source_cash,
         "concept": "Ajuste", "amount": "5.00"},
    ]
    for movement in movements:
        response = client.post("/api/v1/transactions", headers=auth, json={"date": today, **movement})
        assert response.status_code == 201, response.text

    exported = client.get(f"/api/v1/exports/transactions?format={format}", headers=auth)
    assert exported.status_code == 200, exported.text
    if format == "json":
        assert exported.json()[1]["category"] == "Casa personal / Alquiler"
        assert exported.json()[1]["source_account"] == "Banco [EUR]"
        assert all(not any(key.endswith("_id") for key in row) for row in exported.json())

    destination = client.post("/api/v1/auth/register", json={
        "email": "target@example.com", "password": "another-strong-password",
    })
    assert destination.status_code == 201, destination.text
    target_auth = {"Authorization": f"Bearer {destination.json()['access_token']}"}
    account(target_auth, "Otra cuenta")
    category(target_auth, "Otra categoría")
    target_cash = account(target_auth, "Efectivo")
    target_bank = account(target_auth, "Banco")
    target_home = category(target_auth, "Casa personal")
    target_rent = category(target_auth, "Alquiler", target_home)
    assert target_bank != source_bank and target_rent != source_rent

    uploaded = client.post("/api/v1/imports", headers=target_auth, files={
        "file": (f"movimientos.{format}", exported.content),
    })
    assert uploaded.status_code == 201, uploaded.text
    job = uploaded.json()
    assert job["total_rows"] == 4
    assert "source_account" in job["headers"]
    mapping = {field: field for field in job["headers"] if field in {
        "date", "type", "concept", "amount", "source_account", "destination_account",
        "category", "status", "notes", "payment_method", "is_fixed", "is_necessary",
    }}
    preview = client.post(f"/api/v1/imports/{job['id']}/preview", headers=target_auth, json={
        "sheet_name": job["selected_sheet"], "mapping": mapping,
    })
    assert preview.status_code == 200, preview.text
    assert preview.json()["error_rows"] == 0, preview.json()["rows"]
    assert all(row["status"] == "READY" for row in preview.json()["rows"])
    confirmed = client.post(f"/api/v1/imports/{job['id']}/confirm", headers=target_auth)
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["imported_rows"] == 4
    imported = client.get("/api/v1/transactions?page_size=100", headers=target_auth).json()["items"]
    assert len(imported) == 4
    by_concept = {item["concept"]: item for item in imported}
    assert by_concept["Nómina"]["destination_account_id"] == target_bank
    assert by_concept["=Alquiler"]["source_account_id"] == target_bank
    assert by_concept["=Alquiler"]["category_id"] == target_rent
    assert by_concept["=Alquiler"]["status"] == "PENDING"
    assert by_concept["=Alquiler"]["is_fixed"] is True
    assert by_concept["=Alquiler"]["is_necessary"] is True
    assert by_concept["=Alquiler"]["notes"] == "+nota"
    assert by_concept["Traspaso"]["destination_account_id"] == target_cash
    assert by_concept["Ajuste"]["source_account_id"] == target_cash


def test_portable_import_reports_missing_or_ambiguous_accounts(client, auth):
    client.post("/api/v1/accounts", headers=auth, json={
        "name": "Banco", "type": "CHECKING", "currency": "EUR",
    })
    client.post("/api/v1/accounts", headers=auth, json={
        "name": "Banco", "type": "CHECKING", "currency": "EUR",
    })
    content = f"date,concept,amount,type,source_account\n{date.today()},Compra,10,EXPENSE,Banco [EUR]\n"
    uploaded = client.post("/api/v1/imports", headers=auth, files={
        "file": ("movimientos.csv", content.encode()),
    }).json()
    payload = {"sheet_name": "CSV", "mapping": {
        "date": "date", "concept": "concept", "amount": "amount",
        "type": "type", "source_account": "source_account",
    }}
    preview = client.post(f"/api/v1/imports/{uploaded['id']}/preview", headers=auth, json=payload)
    assert preview.json()["rows"][0]["status"] == "ERROR"
    assert "ambigua" in preview.json()["rows"][0]["error"]

    missing = content.replace("Banco [EUR]", "Cuenta inexistente [EUR]")
    uploaded = client.post("/api/v1/imports", headers=auth, files={
        "file": ("movimientos.csv", missing.encode()),
    }).json()
    preview = client.post(f"/api/v1/imports/{uploaded['id']}/preview", headers=auth, json=payload)
    assert preview.json()["rows"][0]["status"] == "ERROR"
    assert "no encontrada" in preview.json()["rows"][0]["error"]


def test_empty_json_export_can_be_uploaded(client, auth):
    exported = client.get("/api/v1/exports/transactions?format=json", headers=auth)
    assert exported.json() == []
    uploaded = client.post("/api/v1/imports", headers=auth, files={
        "file": ("movimientos.json", exported.content),
    })
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["total_rows"] == 0
    assert "source_account" in uploaded.json()["headers"]
